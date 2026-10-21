"""Build the editable assignment report draft from genuine emulator captures."""
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'Report' / '402411354_Reddy_MobileAppDev700_Assignment.docx'
S = ROOT / 'Screenshots'

def font(size=24):
    path = Path('C:/Windows/Fonts/arial.ttf')
    return ImageFont.truetype(str(path), size) if path.exists() else ImageFont.load_default()

def diagram(name, boxes, arrows):
    image = Image.new('RGB', (1100, 450), '#FFF9EE')
    d = ImageDraw.Draw(image)
    for label, rect in boxes.items():
        d.rounded_rectangle(rect, radius=18, fill='white', outline='#287A50', width=3)
        d.multiline_text(((rect[0]+rect[2])/2, (rect[1]+rect[3])/2), label, fill='#26332B', font=font(23), anchor='mm', align='center')
    for start, end in arrows:
        a, b = boxes[start], boxes[end]
        x1, y1 = (a[0]+a[2])/2, a[3]
        x2, y2 = (b[0]+b[2])/2, b[1]
        d.line((x1,y1,x2,y2), fill='#287A50', width=4)
        d.polygon([(x2,y2),(x2-8,y2-13),(x2+8,y2-13)], fill='#287A50')
    path = ROOT / 'Report' / name
    image.save(path)
    return path

screen = diagram('Screen_Flow.png', {
    'Home': (440,20,660,90), 'Pantry': (60,160,280,230),
    'Recipes': (440,160,660,230), 'Settings': (820,160,1040,230),
    'Add/Edit': (60,340,280,410), 'Detail': (440,340,660,410)
}, [('Home','Pantry'),('Home','Recipes'),('Home','Settings'),('Pantry','Add/Edit'),('Recipes','Detail')])

er = diagram('Database_ER.png', {
    'Pantry': (50,70,330,200), 'Recipes': (410,70,690,200),
    'Recipe ingredients': (780,270,1080,400)
}, [('Recipes','Recipe ingredients')])

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.7); sec.bottom_margin = Inches(.65)
sec.left_margin = Inches(.75); sec.right_margin = Inches(.75)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.space_after = Pt(6)
for name, size in [('Title',25),('Heading 1',16),('Heading 2',12)]:
    style=styles[name]; style.font.name='Aptos Display' if name!='Heading 2' else 'Aptos'; style.font.size=Pt(size); style.font.color.rgb=RGBColor(0,0,0)
title_ppr=styles['Title']._element.get_or_add_pPr()
for child in list(title_ppr):
    if child.tag==qn('w:pBdr'): title_ppr.remove(child)
border=OxmlElement('w:pBdr'); bottom=OxmlElement('w:bottom'); bottom.set(qn('w:val'),'nil'); border.append(bottom); title_ppr.append(border)

def h(text, level=1): doc.add_heading(text, level)
def p(text): doc.add_paragraph(text)
def caption(text):
    par=doc.add_paragraph(text); par.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in par.runs: r.italic=True; r.font.size=Pt(9)
def code(text):
    par=doc.add_paragraph()
    par.style='Normal'
    for line in text.splitlines():
        run=par.add_run(line+'\n'); run.font.name='Consolas'; run.font.size=Pt(8)
    return par

title=doc.add_paragraph(style='Title'); title.add_run('Smart Pantry Manager Practical Assignment Report')
p('Mobile App Development 700')
p('Student name: Shiven Reddy')
p('Student number: 402411354')
p('Date: 21 October 2026')
p('GitHub repository: [Add working public repository URL after publication]')
p('This editable draft documents the current Java Android implementation and its verified emulator behavior. Complete the personal reflection, final GitHub link, and any additional evidence before submission.')
doc.add_page_break()

h('Table of contents')
for item in ['1 Introduction .................................. 3','2 System design ................................. 3','3 Screenshots of every output ................... 4','4 Key code snippets ........................... 11','5 Challenges and solutions .................... 12','6 Conclusion and reflection ................... 12','7 Reference list .............................. 12']:
    p(item)
doc.add_page_break()

h('1 Introduction')
p('Food can go to waste when households lose track of what they already own. Smart Pantry Manager records ingredients and suggests only meals whose complete ingredient requirements are present in sufficient quantities. It is intended for people who want to use leftovers before buying more food.')
p('The application is a native Android app written in Java. It stores pantry records and 20 prepared recipes locally in SQLite, works without a network connection, and recalculates suggestions from the current pantry each time the recipes screen resumes.')

h('2 System design')
h('Screen navigation',2)
doc.add_picture(str(screen), width=Inches(6.7)); caption('Figure 1. Activity navigation. The bottom bar connects Home, Pantry, Recipes, and Settings.')
p('Add/Edit receives a pantry item ID through an explicit Intent when editing. Recipe Detail receives a recipe ID. The visible back actions return to the parent screens.')
h('Data model',2)
doc.add_picture(str(er), width=Inches(6.7)); caption('Figure 2. SQLite tables. A recipe has many required ingredient rows; pantry rows are separate user data.')
p('The pantry_items table contains id, name, normalized_name, quantity, unit and optional expiry_date. The recipes table contains id, name, description, instructions and preparation_time. The recipe_ingredients table contains recipe_id, ingredient_name, normalized_name, quantity and unit. A foreign key links each requirement to its recipe. SQLiteOpenHelper creates the schema and seeds the recipes once.')

h('3 Screenshots of every output')
p('The following are genuine captures from an Android emulator running the app on 21 October 2026. They show the states tested during development.')
shots = [
    ('Home.png','Initial Home screen with an empty pantry'),
    ('Pantry_List.png','Pantry list reading a saved ingredient from SQLite'),
    ('Add_Ingredient.png','Ingredient creation form'),
    ('Edit_Ingredient.png','Edit form loaded from a saved record'),
    ('Validation_Error.png','Field-level error for a blank name'),
    ('Delete_Confirmation.png','Confirmation before deleting butter'),
    ('After_Delete.png','Pantry list after deleting butter'),
    ('No_Matching_Recipes.png','No recipe suggested after butter was reduced'),
    ('Suggested_Recipes.png','Buttered toast suggested with both ingredients sufficient'),
    ('Recipe_Detail.png','Full recipe ingredient list and numbered method'),
    ('Settings.png','Persisted expiry indicator setting and About section'),
    ('Reopened_Home.png','Saved pantry count after force-stop and reopening the app'),
    ('Date_Picker.png','Android date picker for the optional expiry date'),
    ('Expiry_Indicator.png','Pantry item marked expiring soon when alerts are enabled'),
]
for i in range(0,len(shots),2):
    table=doc.add_table(rows=2, cols=2); table.alignment=WD_TABLE_ALIGNMENT.CENTER
    for j,(filename,label) in enumerate(shots[i:i+2]):
        path=S/filename
        cell=table.cell(0,j); cell.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].add_run().add_picture(str(path), width=Inches(3.0))
        table.cell(1,j).text=f'Figure {i+j+3}. {label}.'
    doc.add_paragraph('')

doc.add_page_break()
h('4 Key code snippets')
h('SQLite create operation',2)
code('public long addPantryItem(String name, double quantity, String unit, String expiry) {\n    return getWritableDatabase().insertOrThrow(\n        "pantry_items", null, pantryValues(name, quantity, unit, expiry));\n}')
p('DatabaseHelper.java writes validated pantry data to SQLite. The same helper provides read, update and delete operations.')
h('RecyclerView binding',2)
code('PantryItem item = items.get(position);\nholder.name.setText(item.name);\nholder.meta.setText(UnitConverter.format(item.quantity) + " " + item.unit);')
p('PantryAdapter.java binds a dynamic list of SQLite model objects to reusable row views. Edit and delete listeners act on each row’s ID.')
h('Strict matching',2)
code('for (RecipeIngredient need : recipe.ingredients) {\n    String key=need.normalizedName+"|"+UnitConverter.dimension(need.unit);\n    requiredAmounts.put(key,requiredAmounts.getOrDefault(key,0d)\n        +UnitConverter.baseAmount(need.quantity,need.unit));\n}\nfor (Map.Entry<String,Double> requirement:requiredAmounts.entrySet()) {\n    String[] parts=requirement.getKey().split("\\\\|",2);\n    double available=0;\n    for (PantryItem item:pantry)\n        if(parts[0].equals(IngredientNormalizer.normalize(item.name))\n            && parts[1].equals(UnitConverter.dimension(item.unit)))\n            available+=UnitConverter.baseAmount(item.quantity,item.unit);\n    double required=requirement.getValue();\n    if(available+1e-9*Math.max(1d,required)<required) return false;\n}')
p('RecipeMatchingService.java first combines repeated requirements for the same ingredient and dimension, then rejects a recipe if any total is missing or insufficient.')
h('Intent navigation and validation',2)
code('startActivity(new Intent(this, AddEditIngredientActivity.class)\n    .putExtra("item_id", item.id));\nif (cleanName.isEmpty()) { nameLayout.setError("Enter an ingredient name"); return; }')
p('PantryActivity.java passes the record ID to the form. AddEditIngredientActivity.java validates input before any database write.')

doc.add_page_break()
h('5 Challenges and solutions')
p('The first offline Gradle build failed because AndroidX and Material dependencies were not in the local cache. Running an approved online build downloaded those dependencies; subsequent offline builds and unit tests passed.')
p('The emulator exposed a navigation mismatch when Add Ingredient was opened from Home: the visible back text promised Pantry but originally returned Home. The back action was changed to open Pantry explicitly.')
p('A card entrance fade remained visually stale during emulator captures. The adapter fade was removed and normal row opacity restored to keep the list readable.')

h('6 Conclusion and reflection')
p('The current implementation demonstrates Java Activities, explicit Intents, custom RecyclerView adapters, SQLite CRUD, persistent preferences and strict recipe filtering. Automated tests pass for the matching edge cases, and an emulator run verified the core create, edit, delete, restart and recipe appearance/disappearance flow.')
p('[Student: write your own reflection here. Describe what you learned, what you found difficult, and what you would improve with more time.]')

h('7 Reference list')
p('Faculty of Information Technology (2026) Mobile App Development 700 Practical Assignment. Assignment specification supplied for this project.')
p('[Student: add any documentation, tutorials, libraries, or images you actually consulted, using consistent Harvard style.]')

doc.save(OUT)
print(OUT)
